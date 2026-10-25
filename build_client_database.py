"""Pure B2B Client Database Builder for Kazakhstan.
Collects companies across the 6 target niches:
1. E-commerce & Online retail
2. Medical centers & clinics
3. Real estate & developers
4. Legal & accounting firms
5. IT & Digital agencies
6. Manufacturing & construction

Exports directly into Excel and CSV with complete cards (Name, Site, Phone, WhatsApp, Email, City, Niche).
"""

import sys
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from rich.console import Console
from rich.table import Table
from rich.panel import Panel

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from collectors.kz_companies_collector import collect_kz_database

console = Console()
OUTPUT_DIR = Path(__file__).parent / "output"
OUTPUT_DIR.mkdir(exist_ok=True)

EXCEL_FILE = OUTPUT_DIR / "база_клиентов_казахстан.xlsx"
CSV_FILE = OUTPUT_DIR / "база_клиентов_казахстан.csv"


def export_client_database_excel(companies: list, file_path: Path):
    """Exports structured company records into clean, styled Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = "База Клиентов (Казахстан)"

    headers = [
        "№",
        "Компания",
        "Сфера / Ниша",
        "Город",
        "Телефон",
        "WhatsApp",
        "Email",
        "Сайт",
        "Фактический адрес",
        "Источник",
    ]

    # Styles
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")  # Dark slate
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    ws.row_dimensions[1].height = 26

    # Data rows
    for idx, c in enumerate(companies, start=1):
        row_num = idx + 1
        whatsapp_val = c.get("whatsapp", "")
        if whatsapp_val and not whatsapp_val.startswith("http"):
            wa_display = whatsapp_val
        else:
            wa_display = whatsapp_val

        row_data = [
            idx,
            c.get("company", ""),
            c.get("niche", ""),
            c.get("city", "Казахстан"),
            c.get("phone", ""),
            wa_display,
            c.get("email", ""),
            c.get("website", ""),
            c.get("address", ""),
            c.get("source", ""),
        ]
        ws.append(row_data)

        # Alternating light background
        row_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid") if idx % 2 == 0 else PatternFill(fill_type=None)

        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.font = Font(name="Calibri", size=10)
            cell.border = thin_border
            cell.alignment = left_align
            if row_fill.fill_type:
                cell.fill = row_fill

            if col_idx in [1, 4]:
                cell.alignment = center_align

        ws.row_dimensions[row_num].height = 22

    # Column widths
    widths = {
        1: 6,   # №
        2: 32,  # Компания
        3: 28,  # Ниша
        4: 14,  # Город
        5: 22,  # Телефон
        6: 18,  # WhatsApp
        7: 24,  # Email
        8: 28,  # Сайт
        9: 36,  # Адрес
        10: 24, # Источник
    }
    for col_idx, width in widths.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    wb.save(file_path)
    console.print(f"[bold green]✔ Excel база успешно сохранена:[/bold green] {file_path}")


def export_client_database_csv(companies: list, file_path: Path):
    """Exports to CSV in utf-8-sig."""
    df = pd.DataFrame(companies)
    cols = ["company", "niche", "city", "phone", "whatsapp", "email", "website", "address", "source"]
    existing_cols = [c for c in cols if c in df.columns]
    df[existing_cols].to_csv(file_path, index=False, encoding="utf-8-sig")
    console.print(f"[bold green]✔ CSV база успешно сохранена:[/bold green] {file_path}")


def main():
    console.print(
        Panel.fit(
            "[bold cyan]🏢 СБОР БАЗЫ КЛИЕНТОВ (КАЗАХСТАН)[/bold cyan]\n"
            "[dim]Парсинг компаний по 6 целевым сферам с прямыми контактами, сайтами и телефонами[/dim]",
            border_style="cyan",
        )
    )

    companies = collect_kz_database()

    if not companies:
        console.print("[red]Не удалось собрать компании. Проверьте интернет.[/red]")
        return

    # Export
    export_client_database_excel(companies, EXCEL_FILE)
    export_client_database_csv(companies, CSV_FILE)

    # Statistics table
    df = pd.DataFrame(companies)
    niche_counts = df["niche"].value_counts()

    table = Table(title="📊 Статистика собранной базы по отраслям", border_style="cyan")
    table.add_column("Сфера бизнеса", style="bold white")
    table.add_column("Количество компаний", justify="center", style="bold green")

    for niche, count in niche_counts.items():
        table.add_row(str(niche), str(count))

    console.print(table)

    # Sample preview table
    preview_table = Table(title="👀 Пример первых 6 компаний в базе", border_style="dim")
    preview_table.add_column("Компания", style="bold white", width=25)
    preview_table.add_column("Ниша", width=25)
    preview_table.add_column("Город", width=10)
    preview_table.add_column("Телефон / WhatsApp", style="cyan", width=20)
    preview_table.add_column("Сайт", style="blue", width=24)

    for c in companies[:6]:
        contact = c.get("whatsapp") or c.get("phone") or "—"
        site = c.get("website") or "—"
        preview_table.add_row(
            c.get("company", "")[:24],
            c.get("niche", "")[:24],
            c.get("city", "Казахстан")[:10],
            contact[:20],
            site[:24],
        )

    console.print(preview_table)
    console.print(f"\n[bold green]Готово! База из {len(companies)} компаний сформирована.[/bold green]")


if __name__ == "__main__":
    main()

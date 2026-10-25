"""Lead Exporter for RAG B2B Lead Hunter.
Exports enriched leads to styled Excel (.xlsx), CSV (utf-8-sig), and Markdown Digest.
"""

from typing import List, Dict, Any
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def export_to_excel(leads: List[Dict[str, Any]], output_path: Path):
    """Exports leads to a styled Excel sheet with readable column widths and colored headers."""
    if not leads:
        print("[!] Нет данных для экспорта в Excel.")
        return

    # Sort by rag_score descending
    sorted_leads = sorted(leads, key=lambda x: x.get("rag_score", 0), reverse=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "RAG Лиды (Приоритет)"

    headers = [
        "Оценка (1-10)",
        "Компания",
        "Вакансия / Роль",
        "Выявленная боль",
        "Источник",
        "Контакт для связи",
        "Ссылка",
        "Готовый питч (Telegram)",
        "Готовое письмо (Email)",
        "Статус",
    ]

    # Header style
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="top", wrap_text=True)

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    ws.row_dimensions[1].height = 28

    # Score color fills
    high_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")  # soft green
    med_fill = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")   # soft yellow

    for row_idx, lead in enumerate(sorted_leads, start=2):
        score = lead.get("rag_score", 5)
        contact = lead.get("contact", "")
        if contact.startswith("@"):
            contact_str = f"https://t.me/{contact.lstrip('@')}"
        else:
            contact_str = contact

        row_data = [
            score,
            lead.get("company", "Не указана"),
            lead.get("title", ""),
            lead.get("category", ""),
            lead.get("source", ""),
            contact_str,
            lead.get("url", ""),
            lead.get("tg_pitch", ""),
            lead.get("email_pitch", ""),
            "Новый (к отправке)",
        ]
        ws.append(row_data)

        # Style row cells
        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = Font(name="Calibri", size=10)
            cell.border = thin_border
            cell.alignment = left_align

            # Center alignment for specific columns
            if col_idx in [1, 5, 10]:
                cell.alignment = center_align

            # Highlight score
            if col_idx == 1:
                if score >= 8:
                    cell.fill = high_fill
                    cell.font = Font(name="Calibri", size=11, bold=True, color="166534")
                elif score >= 6:
                    cell.fill = med_fill
                    cell.font = Font(name="Calibri", size=10, bold=True, color="854D0E")

        ws.row_dimensions[row_idx].height = 65

    # Column widths
    col_widths = {
        1: 14,  # Score
        2: 24,  # Company
        3: 28,  # Role
        4: 26,  # Pain
        5: 18,  # Source
        6: 25,  # Contact
        7: 30,  # URL
        8: 45,  # TG pitch
        9: 45,  # Email pitch
        10: 18, # Status
    }
    for col_idx, width in col_widths.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    wb.save(output_path)
    print(f"[+] Excel сохранен: {output_path}")


def export_to_csv(leads: List[Dict[str, Any]], output_path: Path):
    """Exports leads to CSV with utf-8-sig encoding for seamless Excel compatibility."""
    if not leads:
        return

    df = pd.DataFrame(leads)
    cols_to_export = [
        "rag_score",
        "company",
        "title",
        "category",
        "source",
        "contact",
        "url",
        "salary",
        "tg_pitch",
        "email_pitch",
    ]
    existing_cols = [c for c in cols_to_export if c in df.columns]
    df_export = df[existing_cols].sort_values(by="rag_score", ascending=False)
    df_export.to_csv(output_path, index=False, encoding="utf-8-sig")
    print(f"[+] CSV сохранен: {output_path}")


def export_markdown_digest(leads: List[Dict[str, Any]], output_path: Path, top_n: int = 10):
    """Generates a high-priority Markdown digest with ready-to-copy pitches."""
    if not leads:
        return

    sorted_leads = sorted(leads, key=lambda x: x.get("rag_score", 0), reverse=True)
    top_leads = sorted_leads[:top_n]

    lines = [
        "# 🎯 Дайджест горячих B2B-лидов для внедрения RAG-систем",
        f"Всего отобрано лидов: **{len(sorted_leads)}**. Ниже — ТОП-{len(top_leads)} с максимальным приоритетом.\n",
    ]

    for idx, lead in enumerate(top_leads, start=1):
        score = lead.get("rag_score", 5)
        company = lead.get("company", "Компания")
        title = lead.get("title", "")
        contact = lead.get("contact", "В вакансии")
        url = lead.get("url", "")
        pain = lead.get("category", "")
        tg_pitch = lead.get("tg_pitch", "")

        contact_md = f"[{contact}](https://t.me/{contact.lstrip('@')})" if contact.startswith("@") else contact

        lines.append(f"### {idx}. {company} — {title}")
        lines.append(f"- **Приоритет:** ⭐ {score}/10")
        lines.append(f"- **Выявленная боль:** {pain}")
        lines.append(f"- **Источник:** {lead.get('source', '')}")
        lines.append(f"- **Контакт:** {contact_md}")
        lines.append(f"- **Ссылка на пост/вакансию:** {url}\n")
        lines.append("**Готовое сообщение для отправки:**")
        lines.append("```text")
        lines.append(tg_pitch)
        lines.append("```\n---\n")

    output_path.write_text("\n".join(lines), encoding="utf-8")
    print(f"[+] Markdown дайджест сохранен: {output_path}")
